# Mini LLM – powered by AI-Implements · C. Lins
# Copyright (c) 2026 C. Lins / AI-Implements – MIT-Lizenz, siehe LICENSE.
# Dieser Code darf frei verwendet, verändert und erweitert werden.
# Dieser Hinweis muss in allen Kopien und abgeleiteten Werken erhalten bleiben.
"""Rechnungs-Skill — beobachtet am 18.10.2026.

Mit aktivem Brief-Skill wurde aus „Schreibe eine Rechnung … 3880€ brutto …
980€ … Iban …“ ein Brief mit „Aktenzeichen: RE/2026/0045“, ohne Netto und
Umsatzsteuer. Die Nachfrage „die Posten als saubere Tabelle mit Summe“ kam
als Strichlinien-Tabelle im Fließtext zurück.
"""
from __future__ import annotations

import datetime as dt
import json
import unittest
import uuid
import zipfile
from io import BytesIO
from types import SimpleNamespace
from unittest.mock import patch

import httpx

import app.main as main
from app.exports import build_export, invoice_email_text
from app.rechnung import (
    betrag,
    euro,
    iban_gueltig,
    invoice_plain_text,
    normalize_invoice_artifact,
)
from app.task_skills import normalize_letter_artifact, resolve_task_skill

GUELTIGE_IBAN = "DE89 3704 0044 0532 0130 00"   # Standard-Beispiel-IBAN
AUFTRAG = (
    "Ok Schreibe ein Rechnung ebenfalls an die oben benannte Person. Leistungen : "
    "implentierung des ai Core Email System - 3880€ brutto, bereitstellen des Mini chat "
    "bot system inkl. Joshi system - inkl., Jahres Nutzungs Gebühr inkl. Support von 10h. "
    "980€, Auflistung gesamt Summe. Innerhalb von 14 Werktagen an das Konto : Iban "
    f"{GUELTIGE_IBAN}. Freundliche Grüße"
)
MODELLANTWORT = {
    "recipient": {"name": "Muster Mann", "department": "", "street": "Musterstraße 123",
                  "postal_code": "12345", "city": "Muster Ort", "country": "Deutschland"},
    "invoice_number": "", "invoice_date": "", "service_date": "", "customer_number": "",
    "subject": "Rechnung für KI-Systeme", "salutation": "Sehr geehrter Herr Mann,",
    "intro": "für die folgenden Leistungen stelle ich Ihnen in Rechnung:",
    "items": [
        {"description": "Implementierung des AI Core E-Mail-Systems", "quantity": 1,
         "unit": "pauschal", "unit_price": 3880, "included": False},
        {"description": "Bereitstellung Mini-Chatbot-System inkl. JOSHI", "quantity": 1,
         "unit": "pauschal", "unit_price": 0, "included": True},
        {"description": "Jahresnutzungsgebühr inkl. 10 Std. Support", "quantity": 1,
         "unit": "Jahr", "unit_price": 980, "included": False},
    ],
    "prices_are": "brutto", "vat_rate": 0, "payment_days": 14, "payment_day_type": "Werktage",
    "iban": GUELTIGE_IBAN, "bic": "", "bank": "", "closing": "Freundliche Grüße", "notes": [],
}
PROFIL = {
    "name": "Max Mustermann", "organization": "Beispiel GmbH", "email": "max@example.com",
    "phone": "", "address": {"street": "Weg 1", "postal_code": "11111", "city": "Musterstadt",
                             "country": "Deutschland"},
    "invoice": {"tax_id": "123/456/78901", "vat_mode": "standard", "iban": "", "bic": "", "bank": ""},
}
JETZT = dt.datetime(2026, 10, 19, 14, 32)


def rechnung(antwort=None, profil=PROFIL, aktiv=True, auftrag=AUFTRAG):
    return normalize_invoice_artifact(antwort or MODELLANTWORT, profil, aktiv, "19.10.2026", auftrag, JETZT)


class RechnenTests(unittest.TestCase):
    def test_gross_prices_are_split_into_net_and_vat(self):
        r = rechnung()
        self.assertEqual(r["gross"], 4860.0)
        self.assertEqual(r["net"], 4084.03)
        self.assertEqual(r["vat"], 775.97)
        self.assertEqual(r["vat_rate"], 19.0)
        self.assertTrue(r["items"][1]["included"])
        self.assertEqual(r["items"][1]["total"], 0.0)

    def test_net_prices_get_vat_on_top_and_quantities_multiply(self):
        antwort = dict(MODELLANTWORT, prices_are="netto", vat_rate=19, items=[
            {"description": "Beratung", "quantity": 2.5, "unit": "Std.", "unit_price": 120, "included": False}])
        r = rechnung(antwort, auftrag="2,5 Std. Beratung à 120 € netto")
        self.assertEqual((r["net"], r["vat"], r["gross"]), (300.0, 57.0, 357.0))
        self.assertEqual(r["items"][0]["total"], 300.0)

    def test_small_business_charges_no_vat_and_says_why(self):
        profil = dict(PROFIL, invoice=dict(PROFIL["invoice"], vat_mode="small"))
        r = rechnung(profil=profil)
        self.assertEqual((r["net"], r["vat"], r["gross"]), (4860.0, 0.0, 4860.0))
        self.assertIn("Gemäß § 19 UStG wird keine Umsatzsteuer berechnet.", r["notes"])

    def test_numbers_are_read_like_germans_write_them(self):
        self.assertEqual(betrag("3.880,00 €"), 3880)
        self.assertEqual(betrag("1,5"), 1.5)
        self.assertEqual(betrag("3880"), 3880)
        self.assertEqual(euro(4084.03), "4.084,03 €")

    def test_a_stored_invoice_is_stable_when_rendered_again(self):
        erste = rechnung()
        zweite = normalize_invoice_artifact(erste)
        for feld in ("invoice_number", "net", "vat", "gross", "payment_terms", "sender", "tax_id", "hints"):
            self.assertEqual(zweite[feld], erste[feld], feld)


class ToleranzTests(unittest.TestCase):
    def test_other_field_names_are_understood(self):
        # glm-5.3-flash, 19.10.2026: „positions“ statt „items“.
        antwort = {k: v for k, v in MODELLANTWORT.items() if k != "items"}
        antwort["positions"] = MODELLANTWORT["items"]
        self.assertEqual(rechnung(antwort)["gross"], 4860.0)
        deutsch = {"empfänger": MODELLANTWORT["recipient"], "positionen": [
            {"bezeichnung": "Beratung", "menge": "2", "einzelpreis": "100,00 €", "inklusive": "nein"},
            {"leistung": "Pflege", "betrag": "50"}], "preisart": "netto", "steuersatz": 19}
        r = rechnung(json.dumps(deutsch), auftrag="Beratung 2 x 100 € netto, Pflege 50 €")
        self.assertEqual([p["total"] for p in r["items"]], [200.0, 50.0])
        self.assertEqual(r["gross"], 297.5)
        self.assertEqual(r["recipient"]["name"], "Muster Mann")

    def test_flat_or_written_out_recipients_are_understood(self):
        flach = {k: v for k, v in MODELLANTWORT.items() if k != "recipient"}
        flach.update({"recipient_name": "Bäckerei Korn", "recipient_street": "Hauptstraße 5",
                      "recipient_zip": "12345", "recipient_city": "Musterstadt", "notes": ""})
        r = rechnung(flach)
        self.assertEqual(r["recipient"]["name"], "Bäckerei Korn")
        self.assertEqual(r["recipient"]["postal_code"], "12345")
        self.assertEqual(r["notes"], [])
        text = dict(MODELLANTWORT, recipient="Bäckerei Korn, Hauptstraße 5, 12345 Musterstadt")
        empfaenger = rechnung(text)["recipient"]
        self.assertEqual((empfaenger["name"], empfaenger["street"], empfaenger["city"]),
                         ("Bäckerei Korn", "Hauptstraße 5", "Musterstadt"))
        self.assertNotIn("Anschrift", " ".join(rechnung(text)["hints"]))
        # glm-5.3-flash, 19.10.2026: Straße, PLZ und Ort in einem Feld.
        zusammen = {k: v for k, v in MODELLANTWORT.items() if k != "recipient"}
        zusammen.update({"recipient_name": "Bäckerei Korn", "recipient_address": "Hauptstraße 5, 12345 Musterstadt"})
        empfaenger = rechnung(zusammen)["recipient"]
        self.assertEqual((empfaenger["street"], empfaenger["postal_code"], empfaenger["city"]),
                         ("Hauptstraße 5", "12345", "Musterstadt"))

    def test_names_dates_cents_and_missing_recipients_are_repaired(self):
        teile = dict(MODELLANTWORT, recipient={"company": "", "first_name": "", "last_name": "Weber",
                                               "street": "Lindenallee 12", "zip": "45678", "city": "Beispielstadt"},
                     service_date="2026-10-14")
        r = rechnung(teile)
        self.assertEqual(r["recipient"]["name"], "Weber")
        self.assertEqual(r["service_date"], "14.10.2026")
        # ministral-3:8b, 19.10.2026: 24000 statt 240 €.
        cent = dict(MODELLANTWORT, prices_are="netto", items=[
            {"description": "Fensterreinigung", "quantity": 1, "unit": "pauschal", "unit_price": 24000, "included": False}])
        r = rechnung(cent, auftrag="Fensterreinigung 240 € netto")
        self.assertEqual(r["items"][0]["total"], 240.0)
        self.assertIn("als 240,00 € gelesen", " ".join(r["hints"]))
        datum = rechnung(dict(MODELLANTWORT, service_date=""), auftrag="Fensterreinigung 240 €, Leistung am 14.10.2026.")
        self.assertEqual(datum["service_date"], "14.10.2026")
        ohne = dict(MODELLANTWORT, recipient={})
        r = rechnung(ohne, auftrag="Erstelle eine Rechnung an die Bäckerei Korn, Hauptstraße 5, 12345 Musterstadt: Pflege 60 €")
        self.assertEqual((r["recipient"]["name"], r["recipient"]["postal_code"], r["recipient"]["city"]),
                         ("Bäckerei Korn", "12345", "Musterstadt"))

    def test_a_change_keeps_the_invoice_number(self):
        from app.rechnung import vorige_rechnungsnummer
        alt = invoice_plain_text(rechnung())
        self.assertEqual(vorige_rechnungsnummer(["Hallo", alt]), "RE-20261019-1432")
        neu = normalize_invoice_artifact(MODELLANTWORT, PROFIL, True, "20.10.2026", AUFTRAG,
                                         dt.datetime(2026, 10, 20, 9, 0), vorige_nummer="RE-20261019-1432")
        self.assertEqual(neu["invoice_number"], "RE-20261019-1432")
        self.assertNotIn("automatisch vergeben", " ".join(neu["hints"]))


class PflichtangabenTests(unittest.TestCase):
    def test_number_date_terms_and_sender_are_set_not_invented(self):
        r = rechnung()
        self.assertEqual(r["invoice_number"], "RE-20261019-1432")
        self.assertEqual(r["invoice_date"], "19.10.2026")
        self.assertIn("14 Werktagen", r["payment_terms"])
        self.assertEqual(r["sender"]["name"], "Max Mustermann")
        self.assertEqual(r["tax_id"], "123/456/78901")
        self.assertEqual(r["closing"], "Freundliche Grüße")
        self.assertEqual(r["iban"], GUELTIGE_IBAN)

    def test_missing_mandatory_data_becomes_a_hint(self):
        r = rechnung(profil={}, aktiv=False)
        hinweise = " ".join(r["hints"])
        self.assertIn("Steuernummer", hinweise)
        self.assertIn("Absender fehlt", hinweise)
        self.assertEqual(r["tax_id"], "")

    def test_a_wrong_iban_and_an_invented_amount_are_flagged(self):
        self.assertTrue(iban_gueltig(GUELTIGE_IBAN))
        self.assertFalse(iban_gueltig("DE00 1234 5678 9012 3456 78"))
        antwort = dict(MODELLANTWORT, iban="DE00 1234 5678 9012 3456 78", items=[
            *MODELLANTWORT["items"],
            {"description": "Einrichtung", "quantity": 1, "unit": "pauschal", "unit_price": 450, "included": False}])
        hinweise = " ".join(rechnung(antwort, auftrag=AUFTRAG.replace(GUELTIGE_IBAN, "x"))["hints"])
        self.assertIn("ungültige Prüfziffer", hinweise)
        self.assertIn("450,00 € steht so nicht im Auftrag", hinweise)
        self.assertNotIn("3.880,00 € steht", hinweise)

    def test_calendar_days_get_a_due_date(self):
        antwort = dict(MODELLANTWORT, payment_days=30, payment_day_type="Tage")
        self.assertIn("bis zum 18.11.2026", rechnung(antwort)["payment_terms"])


class ErkennungTests(unittest.TestCase):
    def test_an_invoice_wins_over_an_active_letter_skill(self):
        self.assertEqual(resolve_task_skill(AUFTRAG, ["letter"]), "invoice")
        self.assertEqual(resolve_task_skill("Erstelle eine Rechnung über 500 Euro an Firma Y", []), "invoice")
        self.assertEqual(resolve_task_skill("irgendwas", ["invoice"]), "invoice")

    def test_a_follow_up_keeps_editing_the_invoice(self):
        verlauf = [{"role": "user", "content": AUFTRAG},
                   {"role": "assistant", "content": invoice_plain_text(rechnung())}]
        frage = "Kannst du die posten innerhalb einer sauberen Tabelle anfertigen und unten als Abschluss Summe."
        self.assertEqual(resolve_task_skill(frage, ["letter"], verlauf), "invoice")
        self.assertIsNone(resolve_task_skill("Was bedeutet Leistungsdatum?", [], verlauf))

    def test_calculations_and_questions_are_not_invoices(self):
        self.assertIsNone(resolve_task_skill("mach mal eine Rechnung, wie viel ich im Jahr spare", []))
        self.assertIsNone(resolve_task_skill("Was gehört alles in eine Rechnung?", []))
        self.assertEqual(resolve_task_skill("Schreibe einen Brief an Herrn X", []), "letter")


class BriefGrussTests(unittest.TestCase):
    def test_a_requested_closing_survives_the_letter_skill(self):
        # 18.10.2026: „Freundliche Grüße“ wurde wieder zu „Mit freundlichen Grüßen“.
        brief = normalize_letter_artifact({"subject": "x", "paragraphs": ["Text."], "closing": "Freundliche Grüße"})
        self.assertEqual(brief["closing"], "Freundliche Grüße")
        self.assertEqual(normalize_letter_artifact({"paragraphs": ["a"], "closing": "Tschüss"})["closing"],
                         "Mit freundlichen Grüßen")


class ExportTests(unittest.TestCase):
    def test_every_format_carries_the_same_numbers(self):
        r = rechnung()
        inhalt = invoice_plain_text(r)
        pdf, art, name = build_export("pdf", "x", inhalt, r)
        self.assertTrue(pdf.startswith(b"%PDF"))
        self.assertTrue(name.lower().startswith("rechnung-re-20261019-1432"), name)
        docx, _, _ = build_export("docx", "x", inhalt, r)
        with zipfile.ZipFile(BytesIO(docx)) as archiv:
            dokument = archiv.read("word/document.xml").decode()
        for text in ("4.860,00 €", "775,97 €", "4.084,03 €", "inklusive", "123/456/78901"):
            self.assertIn(text, dokument)
        xlsx, _, _ = build_export("xlsx", "x", inhalt, r)
        from openpyxl import load_workbook
        blatt = load_workbook(BytesIO(xlsx)).active
        werte = [zelle.value for zeile in blatt.iter_rows() for zelle in zeile]
        self.assertIn(4860.0, werte)
        self.assertIn(3880.0, werte)
        eml, _, _ = build_export("eml", "x", inhalt, r)
        self.assertIn("Rechnung RE-20261019-1432", eml.decode("utf-8", "ignore"))
        text = invoice_email_text(r)
        self.assertIn("Gesamtbetrag (brutto): 4.860,00 €", text)
        self.assertNotIn("Vor dem Versand", text, "Hinweise gehören nicht in den Versand.")

    def test_plain_text_has_a_real_table(self):
        text = invoice_plain_text(rechnung())
        self.assertTrue(text.startswith("Rechnung Nr. RE-20261019-1432"))
        self.assertIn("| Pos. | Leistung | Menge | Einzelpreis | Gesamt |", text)
        self.assertNotIn("------------", text)


class EndpunktTests(unittest.IsolatedAsyncioTestCase):
    async def test_the_chat_returns_a_computed_invoice(self):
        gesendet = []

        def ollama(anfrage: httpx.Request) -> httpx.Response:
            nutzlast = json.loads(anfrage.content)
            gesendet.append(nutzlast)
            return httpx.Response(200, json={"message": {"content": json.dumps(MODELLANTWORT)}, "done": True})

        class Auftraege:
            async def start(self, job, strom):
                self.strom = strom

            async def subscribe(self, request_id, user_id):
                async for zeile in self.strom:
                    yield zeile

        async with httpx.AsyncClient(transport=httpx.MockTransport(ollama)) as client:
            anfrage = SimpleNamespace(app=SimpleNamespace(state=SimpleNamespace(http=client)))
            with patch.object(main, "require_user", return_value={"id": 1}), \
                    patch.object(main, "chat_jobs", Auftraege()), \
                    patch.object(main, "load_profile", return_value=(True, PROFIL)), \
                    patch.object(main, "allowed_context_ids", return_value=[]):
                antwort = await main.chat(
                    request=anfrage, model="glm-5.3-flash:cloud", prompt=AUFTRAG,
                    history="[]", request_id=f"test-{uuid.uuid4()}", chat_id="c1",
                    user_message_id="u1", assistant_message_id="a1", response_prefix="",
                    previous_usage="{}", temperature=0.7, web_enabled=False,
                    profile_context_enabled=True, persona_mode="global", chat_persona="",
                    active_skills='["letter"]', skill_options="{}", skill_action=False,
                    context_ids="[]", files=None,
                )
                ereignisse = []
                async for zeile in antwort.body_iterator:
                    text = zeile.decode() if isinstance(zeile, bytes) else zeile
                    ereignisse += [json.loads(teil) for teil in text.splitlines() if teil.strip()]
        artefakt = next(e for e in ereignisse if e.get("type") == "artifact")
        self.assertEqual(artefakt["artifact"]["type"], "invoice")
        self.assertEqual(artefakt["artifact"]["gross"], 4860.0)
        self.assertEqual(artefakt["artifact"]["tax_id"], "123/456/78901")
        self.assertTrue(artefakt["content"].startswith("Rechnung Nr. "))
        self.assertEqual(gesendet[0]["format"]["required"][0], "recipient")
        # Steuernummer und IBAN aus dem Profil gehen nicht an das Modell.
        alles = json.dumps(gesendet[0]["messages"], ensure_ascii=False)
        self.assertNotIn("123/456/78901", alles)


class OberflaechenTests(unittest.TestCase):
    def test_the_skill_card_and_profile_fields_exist(self):
        from pathlib import Path
        wurzel = Path(__file__).resolve().parent.parent
        js = (wurzel / "static/app.js").read_text(encoding="utf-8")
        html = (wurzel / "static/index.html").read_text(encoding="utf-8")
        self.assertIn('id: "invoice"', js)
        self.assertIn("function invoiceArtifactHtml", js)
        self.assertIn('message?.artifact?.type === "invoice"', js)
        for feld in ("profile-tax-id", "profile-vat-mode", "profile-iban", "profile-bic", "profile-bank"):
            self.assertIn(f'id="{feld}"', html)


if __name__ == "__main__":
    unittest.main()
